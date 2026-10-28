"""合成內容的離線格式完整性；不呼叫模型。"""
from base64 import b64encode
from io import BytesIO
import json
import sqlite3
import wave
import unicodedata
import re
import zipfile

import av
import pymupdf
import pytest
from docx import Document

from runtime import export_documents as documents, export_map, export_podcast, export_cards
from runtime.api.export_routes import CardImage


def note():
    evidence = {'page': 13, 'normalized_page': 3, 'source_name': '錄音.MP3', 'block_order': 2}
    p = {'id': 'p', 'heading': '繁體中文粗體標題', 'text': '完整條件：x < 3，且不可省略。',
         'formula': 'a = b + 1', 'table': {'headers': ['條件', '結果'], 'rows': [['甲', '乙']]}, 'evidence': [evidence]}
    return {'material_id': 'm', 'material_name': '測試教材', 'knowledge_structure_revision': 'revision', 'kind': 'summary',
            'source_status': {'quality': 'accepted', 'processing': 'succeeded', 'decision': 'retain'},
            'paragraphs': [p], 'annotations': [
                {'paragraph_id': 'p', 'body': '我的想法'},
                {'paragraph_id': None, 'body': '舊註記', 'original_revision': 'old', 'original_paragraph': {**p, 'heading': '原段落'}}]}


def test_notes_preserve_body_table_formula_sources_and_orphaned_annotations():
    view = note()
    markdown = documents.note_markdown(view)
    assert all(text in markdown for text in ('我的想法', '舊註記', '原段落', '第 3 段', '區塊 2', '| 甲 | 乙 |', 'old'))
    word = Document(BytesIO(documents.note_docx(view)))
    assert word.tables[0].cell(1, 0).text == '甲'
    assert '章節大綱' in [p.text for p in word.paragraphs]
    assert '舊註記' in '\n'.join(p.text for p in word.paragraphs)
    with pymupdf.open(stream=documents.note_pdf(view), filetype='pdf') as pdf:
        assert len(pdf) >= 1
        text = unicodedata.normalize('NFKC', ''.join(page.get_text() for page in pdf))
        assert all(t in text for t in ('我的想法', '舊註記', '原段落', '第 3 段', 'a = b + 1'))
        spans = [s for b in pdf[0].get_text('dict')['blocks'] for l in b.get('lines', []) for s in l['spans']]
        assert any('Bold' in s['font'] and '繁體中文' in s['text'] for s in spans)
        assert not any('Droid' in s['font'] and s['size'] == 14 for s in spans)


@pytest.mark.parametrize('kind', ['summary', 'full'])
@pytest.mark.parametrize('structure', ['legacy', 'bullets'])
def test_note_pdf_keeps_heading_with_body_start_and_preserves_long_body(kind, structure):
    view = note()
    view['kind'] = kind
    first = view['paragraphs'][0]
    lines = [f'長正文第{i:03d}行' for i in range(100)]
    view['paragraphs'] = [
        {**first, 'id': 'lead', 'heading': '前文', 'text': '填充行\n' * 40, 'formula': None, 'table': None},
        {**first, 'id': 'target', 'heading': '目標標題 TARGET',
         'text': '首段正文 BODY\n' + '\n'.join(lines), 'formula': None, 'table': None},
    ]
    if structure == 'bullets':
        view['paragraphs'] = [
            {k: v for k, v in paragraph.items() if k not in {'text', 'formula', 'table'}} | {
                'level': 1 if index == 0 else 2,
                'blocks': ([{'type': 'text', 'text': paragraph['text']}] if index == 0 else
                    [{'type': 'bullets', 'items': [{'text': line, 'children': []}
                        for line in paragraph['text'].splitlines()]}])}
            for index, paragraph in enumerate(view['paragraphs'])]
    view['annotations'] = []
    with pymupdf.open(stream=documents.note_pdf(view), filetype='pdf') as pdf:
        pages = [unicodedata.normalize('NFKC', page.get_text()) for page in pdf]
        heading_size = 13 if structure == 'bullets' else 14
        heading_page = next(i for i, page in enumerate(pdf)
            if any('TARGET' in span['text'] and round(span['size']) == heading_size
                for block in page.get_text('dict')['blocks'] for line in block.get('lines', []) for span in line['spans']))
        body_page = next(i for i, text in enumerate(pages) if '首段正文 BODY' in text)
        assert heading_page == body_page
        assert body_page > 0
        assert next(i for i, text in enumerate(pages) if lines[-1] in text) > body_page
        assert all(''.join(pages).count(line) == 1 for line in lines)


def structured_note():
    evidence = {'page': 3, 'normalized_page': 3, 'source_name': '經濟.pdf', 'block_order': 2}
    items = [{'text': '出口占比', 'children': [{'text': '子項', 'children': [{'text': '孫項'}]}]}]
    chapter = {'id': '0', 'heading': '國際經濟學概論', 'level': 1, 'evidence': [evidence],
               'blocks': [{'type': 'bullets', 'items': items}]}
    section = {'id': '1', 'heading': '開放度', 'level': 2, 'evidence': [evidence], 'blocks': [
        {'type': 'formula', 'latex': r'\text{開放度} = \frac{\text{出口}+\text{進口}}{\text{GDP}} \times 100\%'},
        {'type': 'table', 'caption': '各國比較', 'headers': ['國家', '值'], 'rows': [['甲', '50']]},
        {'type': 'table', 'caption': None, 'headers': ['年', '值'], 'rows': [['2020', '1']]},
        {'type': 'callout', 'style': 'case_study', 'title': '甲國', 'items': [{'text': '案例內容', 'children': []}]},
        {'type': 'callout', 'style': 'lesson', 'title': None, 'items': [{'text': '教訓內容', 'children': []}]}]}
    view = note()
    view['paragraphs'] = [chapter, section]
    view['annotations'] = [{'paragraph_id': '1', 'body': '我的想法'},
        {'paragraph_id': None, 'body': '舊註記', 'original_revision': 'old', 'original_paragraph': note()['paragraphs'][0]}]
    return view


def test_structured_notes_keep_outline_levels_bullets_tables_formula_and_callouts_in_every_format():
    view = structured_note()
    markdown = documents.note_markdown(view)
    assert '## 章節大綱' in markdown
    links = re.findall(r'- \[.*?\]\(#([^)]*)\)', markdown)
    assert len(links) == 2 and all(f'<a id="{anchor}"></a>' in markdown for anchor in links)
    assert '## 一、國際經濟學概論' in markdown and '### 1.1 開放度' in markdown
    assert '- 出口占比\n  - 子項\n    - 孫項' in markdown
    assert '$$\n' + view['paragraphs'][1]['blocks'][0]['latex'] + '\n$$' in markdown
    assert '> **案例研究：甲國**' in markdown and '> **重要教訓**' in markdown and '| 2020 | 1 |' in markdown
    assert '原段落' not in markdown and '繁體中文粗體標題' in markdown and '舊註記' in markdown
    word = Document(BytesIO(documents.note_docx(view)))
    texts = [p.text for p in word.paragraphs]
    assert '一、國際經濟學概論' in texts and '1.1 開放度' in texts
    assert [p.style.name for p in word.paragraphs if p.text in ('出口占比', '子項', '孫項')] == ['List Bullet', 'List Bullet 2', 'List Bullet 3']
    assert len(word.inline_shapes) == 1
    assert word.inline_shapes[0]._inline.docPr.get('descr') == view['paragraphs'][1]['blocks'][0]['latex']
    assert [t.cell(1, 0).text for t in word.tables[:2]] == ['甲', '2020']
    assert word.tables[2].cell(0, 0).paragraphs[0].text == '案例研究：甲國'
    with pymupdf.open(stream=documents.note_pdf(view), filetype='pdf') as pdf:
        text = unicodedata.normalize('NFKC', ''.join(page.get_text() for page in pdf))
        assert all(t in text for t in ('章節大綱', '一、國際經濟學概論', '1.1 開放度', '孫項', '各國比較', '案例研究:甲國', '重要教訓', '我的想法', '舊註記'))
        assert len(pdf[0].get_images()) == 1


def test_docx_long_formula_fits_page_width():
    view = structured_note()
    latex = ' + '.join([r'x_{1}'] * 40)
    view['paragraphs'][1]['blocks'][0]['latex'] = latex
    word = Document(BytesIO(documents.note_docx(view)))
    section = word.sections[-1]
    assert word.inline_shapes[0].width <= section.page_width - section.left_margin - section.right_margin
    assert word.inline_shapes[0]._inline.docPr.get('descr') == latex


def test_diagnosis_overflow_is_explicit_and_stays_one_page():
    point = {'title': '需要複習的內容' * 8, 'correct_count': 1, 'answer_count': 3, 'evidence': [{'page': 2, 'source_name': '課程.youtube'}]}
    view = {'scope': 'material', 'strengths': [], 'weaknesses': [{'label': '概念', 'points': [point] * 80}],
            'recommendations': [{'label': '概念', 'reason': '80 個重點尚待加強。'}]}
    with pymupdf.open(stream=documents.diagnosis_pdf(view), filetype='pdf') as pdf:
        assert len(pdf) == 1
        assert '未列於摘要' in unicodedata.normalize('NFKC', pdf[0].get_text()) and '80' in pdf[0].get_text()
        assert '第 2 段' in unicodedata.normalize('NFKC', pdf[0].get_text())


def test_mp3_is_decodable_and_transcript_uses_saved_names_and_time():
    source = BytesIO()
    with wave.open(source, 'wb') as wav:
        wav.setnchannels(1); wav.setsampwidth(2); wav.setframerate(24000)
        wav.writeframes(b'\x00\x01' * 24000)
    source.seek(0)
    mp3 = export_podcast.mp3(source)
    with av.open(BytesIO(mp3)) as audio:
        frames = list(audio.decode(audio=0))
        assert sum(f.samples for f in frames) >= 44000
    episode = {'hosts': {'host': '小安', 'guest': '小禾'}, 'claims': [{'label': '主題', 'text': '本集重點', 'evidence': [{'page': 9, 'normalized_page': 4, 'source_name': '錄音.wav'}]}]}
    text = export_podcast.transcript({'name': '節目', 'episodes': [episode]}, 0, {'segments': [
        {'start': 12, 'end': 24, 'turns': [{'speaker': 'host', 'text': '甲'}, {'speaker': 'guest', 'text': '乙'}]}]})
    assert all(s in text for s in ('小安', '小禾', '00:00:12', '00:00:24', '本集重點', '第 4 段'))


def test_map_covers_all_regions_and_cross_region_relations():
    view = {'material_id': 'm', 'material_name': '測試教材', 'knowledge_structure_revision': 'r',
            'concepts': [{'concept_id': str(i), 'label': f'概念{i}', 'claim_count': 1, 'claims': [{'text': f'完整重點{i}', 'evidence': [{'page':i+1,'source_name':'原文.pdf'}]}]} for i in range(30)],
            'relations': [{'source_concept_id': '0', 'target_concept_id': '29', 'type': 'prerequisite'}]}
    regions = {'regions': [{'label': '區域一', 'concept_ids': [str(i) for i in range(15)]},
                           {'label': '區域二', 'concept_ids': [str(i) for i in range(15, 30)]}]}
    md = export_map.markdown(view, regions)
    assert all(f'概念{i}' in md and f'完整重點{i}' in md for i in range(30))
    assert md.count('```mermaid') == 2 and '相關：先備 → 概念29' in md and '來源：原文.pdf 第 30 頁' in md


def test_card_package_has_all_images_full_text_sources_and_stable_ids(monkeypatch, tmp_path):
    doc = pymupdf.open(); page = doc.new_page(width=120, height=76)
    png = page.get_pixmap().tobytes('png'); doc.close()
    cards = [{'concept_id': str(i), 'label': f'概念{i}', 'claims': [{'text': f'完整背面{i}', 'evidence': [{'page': i + 1, 'source_name': '錄音.mp3'}]}]} for i in range(9)]
    def read(owner, identity, *, offset=0, expected_version=None, dsn=None):
        if expected_version is not None:
            return {'cards': cards[offset:offset + 8]}
        return {'version': 1, 'name': '整組', 'material_name': '教材.pptx', 'navigation': cards, 'batch': {}}
    monkeypatch.setattr(export_cards.card_sets, 'read_card_study', read)
    monkeypatch.setattr(export_cards, '_root', lambda: tmp_path)
    images = [CardImage(concept_id=c['concept_id'], page=p, png=b64encode(png).decode()) for c in cards for p in range(2)]
    def exported(format):
        upload_id = export_cards.create('owner', 'deck', 1, format)['upload_id']
        for offset in range(0, len(images), 2):
            export_cards.append('owner', 'deck', upload_id, offset, images[offset:offset + 2])
        with export_cards.finish('owner', 'deck', upload_id) as (path, *_):
            if format == 'apkg':
                assert (path.parent / 'collection.anki2').is_file()
            content = path.read_bytes()
        assert not (export_cards.work_root() / f'owner_{upload_id}').exists()
        return content
    with monkeypatch.context() as guard:
        def outside_temp(*args, **kwargs):
            raise AssertionError('Anki private SQLite must stay inside the upload workspace')
        guard.setattr('tempfile.mkstemp', outside_temp)
        package = exported('apkg')
    with zipfile.ZipFile(BytesIO(package)) as archive:
        assert len(json.loads(archive.read('media'))) == 18
        path = tmp_path / 'anki.db'; path.write_bytes(archive.read('collection.anki2'))
    with sqlite3.connect(path) as db:
        notes = db.execute('select guid, flds from notes').fetchall()
        assert len(notes) == 9
        assert all('完整背面' in fields and ' 段' in fields and '<img src=' in fields for _, fields in notes)
    with zipfile.ZipFile(BytesIO(exported('zip'))) as archive:
        assert len(archive.namelist()) == 18 and archive.testzip() is None
        assert archive.namelist() == [f"概念{i}-{page}.png" for i in range(9) for page in (1, 2)]
    with pymupdf.open(stream=exported('pdf'), filetype='pdf') as pdf:
        assert len(pdf) == 9
        assert all(len(page.get_images()) == 2 for page in pdf)
    cards[0]['label'] = cards[1]['label'] = '相同/概念?'
    with zipfile.ZipFile(BytesIO(exported('zip'))) as archive:
        names = archive.namelist()
        assert len(names) == len(set(names)) == 18
        assert all('/' not in name and '?' not in name for name in names)
        assert all(name.startswith('相同概念') for name in names[:4])


def test_long_summary_pdf_flows_onto_more_pages_without_shrinking_or_dropping_text():
    view = note()
    first = view['paragraphs'][0]
    view['paragraphs'] = [{**first, 'id': str(i), 'heading': f'繁體中文標題 {i}',
                           'text': f'第 {i} 段完整正文，不裁切也不縮字。' * 40} for i in range(12)]
    view['annotations'] = []
    with pymupdf.open(stream=documents.note_pdf(view), filetype='pdf') as pdf:
        body = [page for page in pdf if '來源與個人註記' not in page.get_text()]
        assert len(body) >= 3 and all(page.rect == pymupdf.paper_rect('a4') for page in body)
        text = unicodedata.normalize('NFKC', ''.join(page.get_text() for page in pdf))
        assert all(f'繁體中文標題 {i}' in text and f'第 {i} 段完整正文' in text for i in range(12))
        spans = [s for page in body for b in page.get_text('dict')['blocks'] for l in b.get('lines', []) for s in l['spans']]
        assert min(s['size'] for s in spans if "正文" in s['text'] or "標題" in s['text']) >= 12
        assert any('Bold' in s['font'] and '繁體中文' in s['text'] for s in spans)


@pytest.mark.parametrize('name', ['來源.mp3', '來源.M4A', '來源.wav', '來源.mp4', '來源.youtube'])
def test_media_sources_use_transcript_segments(name):
    assert documents.source_text({'source_name': name, 'page': 14, 'normalized_page': 4}) == f'{name} · 第 4 段'
    assert documents.source_text({'source_name': '來源.pdf', 'page': 14, 'normalized_page': 4}) == '來源.pdf · 第 4 頁'


def test_docx_cleans_xml_controls_without_changing_the_note_and_sets_author():
    view = note()
    p = view['paragraphs'][0]
    p['heading'] = '標題\x0c'; p['text'] = '正文\x00\x0b\x0c\x1f\ud800\uffff\t保留\n換行'
    p['formula'] = 'x = 1\x0c'
    p['table']['headers'][0] += '\x01'
    p['table']['rows'][0][0] += '\x0c'
    p['evidence'][0]['source_name'] = '錄音\x0c.mp3'
    view['annotations'][0]['body'] = '註記\x0c'
    doc = Document(BytesIO(documents.note_docx(view)))
    assert '正文\t保留\n換行' in [p.text for p in doc.paragraphs]
    assert doc.tables[0].cell(1, 0).text == '甲'
    assert '\x0c' in view['paragraphs'][0]['text']
    assert doc.core_properties.author == doc.core_properties.last_modified_by == 'Studydy'


def test_empty_diagnosis_has_no_empty_recommendations_or_zero_omissions():
    view = {'scope': 'material', 'strengths': [], 'weaknesses': [], 'recommendations': []}
    with pymupdf.open(stream=documents.diagnosis_pdf(view), filetype='pdf') as doc:
        text = unicodedata.normalize('NFKC', doc[0].get_text())
        assert '建議先複習' not in text and '另有' not in text
        assert len(doc) == 1


def test_readable_math_in_bullets_and_cells_survives_all_exports():
    view = structured_note()
    view['paragraphs'][0]['blocks'][0]['items'][0]['text'] = '平方根 √4 = 2'
    view['paragraphs'][1]['blocks'][1]['rows'] = [['甲', '≈ 50']]
    markdown = documents.note_markdown(view)
    assert '平方根 √4 = 2' in markdown and '| 甲 | ≈ 50 |' in markdown
    word = Document(BytesIO(documents.note_docx(view)))
    assert '平方根 √4 = 2' in [p.text for p in word.paragraphs]
    assert word.tables[0].cell(1, 1).text == '≈ 50'
    assert len(word.inline_shapes) == 1
    with pymupdf.open(stream=documents.note_pdf(view), filetype='pdf') as pdf:
        text = unicodedata.normalize('NFKC', ''.join(page.get_text() for page in pdf))
        assert '√4 = 2' in text and '≈ 50' in text
        assert sum(len(page.get_images()) for page in pdf) >= 1


def test_pdf_formulas_keep_docx_size_across_page_boundaries():
    view = structured_note()
    formula = {'type': 'formula', 'latex': r'\frac{x_{1}+1}{y_{2}} = 2'}
    view['paragraphs'] = [{**view['paragraphs'][0], 'blocks': [formula] * 50}]
    view['annotations'] = []
    word = Document(BytesIO(documents.note_docx(view)))
    with pymupdf.open(stream=documents.note_pdf(view), filetype='pdf') as pdf:
        images = [pymupdf.Rect(image['bbox']) for page in pdf for image in page.get_image_info()]
        assert len(images) == len(word.inline_shapes) == 50
        for image, shape in zip(images, word.inline_shapes):
            assert image.width == pytest.approx(shape.width.pt, abs=0.02)
            assert image.height == pytest.approx(shape.height.pt, abs=0.02)


def test_export_metadata_uses_material_name_and_short_hash_without_pdf_clipping():
    view = note()
    view.update(material_id='12345678-1234-1234-1234-123456789abc', material_name='網路協定教材',
                knowledge_structure_revision='knowledge-structure:sha256:' + 'abcdefab' * 8)
    view['annotations'][1]['original_revision'] = 'knowledge-structure:sha256:' + 'fedcbabc' * 8
    markdown = documents.note_markdown(view)
    word = Document(BytesIO(documents.note_docx(view)))
    with pymupdf.open(stream=documents.note_pdf(view), filetype='pdf') as pdf:
        pdf_text = unicodedata.normalize('NFKC', ''.join(page.get_text() for page in pdf))
    for output in (markdown, '\n'.join(p.text for p in word.paragraphs), pdf_text):
        assert all(value in output for value in ('網路協定教材', '12345678', 'abcdefab', 'fedcbabc'))
        assert view['material_id'] not in output and view['knowledge_structure_revision'] not in output
    map_text = export_map.markdown({**view, 'concepts': [], 'relations': []}, {'regions': []})
    assert all(value in map_text for value in ('網路協定教材', '12345678', 'abcdefab'))


def test_markdown_groups_sources_without_losing_file_page_or_block_locations():
    view = note()
    view['annotations'] = []
    first = {'source_id': 'source-a', 'source_name': '甲.pdf', 'page': 31, 'normalized_page': 2}
    evidence = [{**first, 'block_order': i} for i in range(22)]
    evidence += [{**first, 'block_order': 24}, {**first, 'block_order': 24},
                 {**first, 'normalized_page': 3, 'block_order': 0},
                 {**first, 'source_id': 'source-b', 'source_name': '乙.pdf', 'block_order': 1}]
    view['paragraphs'][0]['evidence'] = evidence
    markdown = documents.note_markdown(view)
    sources = [line for line in markdown.splitlines() if line.startswith('來源：')]
    assert sources == ['來源：甲.pdf · 第 2 頁 · 區塊 0–21、24  ',
                       '來源：甲.pdf · 第 3 頁 · 區塊 0  ',
                       '來源：乙.pdf · 第 2 頁 · 區塊 1  ']
    assert len(view['paragraphs'][0]['evidence']) == 26


def test_export_times_use_taipei_for_zip_pdf_and_word(monkeypatch):
    from datetime import datetime, timezone
    from xml.etree import ElementTree as ET
    class FixedTime(datetime):
        @classmethod
        def now(cls, tz=None):
            return datetime(2026, 10, 27, 0, 15, 30, tzinfo=timezone.utc).astimezone(tz)
    monkeypatch.setattr(documents, 'datetime', FixedTime)
    with zipfile.ZipFile(BytesIO(documents.zip_files([('來源.txt', b'content')]))) as archive:
        assert archive.infolist()[0].date_time == (2026, 10, 27, 8, 15, 30)
    with pymupdf.open(stream=documents.pdf('<p>合成內容</p>'), filetype='pdf') as doc:
        assert doc.metadata['creationDate'] == "D:20261027081530+08'00'"
    with zipfile.ZipFile(BytesIO(documents.note_docx(note()))) as archive:
        assert all(item.date_time == (2026, 10, 27, 8, 15, 30) for item in archive.infolist())
        root = ET.fromstring(archive.read('docProps/core.xml'))
        assert root.find('{http://purl.org/dc/terms/}created').text == '2026-10-27T08:15:30+08:00'


@pytest.mark.parametrize("structured", [False, True])
def test_note_markdown_keeps_filenames_and_hyphenated_terms_readable(structured):
    view = structured_note() if structured else note()
    name = "01_生成式AI基礎與機率建模.pptx"
    view["material_name"] = name
    paragraph = view["paragraphs"][0]
    paragraph["evidence"][0]["source_name"] = name
    text = "Self-attention 使用 QK^T；保留 *字面星號* 與 [引用]。"
    if structured:
        paragraph["blocks"] = [{"type": "text", "text": text}]
    else:
        paragraph["text"] = text
    markdown = documents.note_markdown(view)
    assert name in markdown and "Self-attention 使用 QK^T" in markdown
    assert r"\*字面星號\*" in markdown and r"\[引用\]" in markdown


def test_note_word_sources_are_real_footnotes_with_the_pdf_source_grouping():
    from xml.etree import ElementTree as ET
    view = note()
    view["paragraphs"][0]["evidence"] = [*view["paragraphs"][0]["evidence"],
        {"page": 13, "normalized_page": 3, "source_name": "錄音.MP3", "block_order": 3},
        {"page": 4, "normalized_page": 4, "source_name": "課堂.pdf", "block_order": 1},
    ]
    data = documents.note_docx(view)
    assert "我的想法" in "\n".join(paragraph.text for paragraph in Document(BytesIO(data)).paragraphs)
    namespaces = {"w": "http://schemas.openxmlformats.org/wordprocessingml/2006/main"}
    with zipfile.ZipFile(BytesIO(data)) as archive:
        body = ET.fromstring(archive.read("word/document.xml"))
        notes = ET.fromstring(archive.read("word/footnotes.xml"))
        references = body.findall(".//w:footnoteReference", namespaces)
        assert [reference.attrib[f"{{{namespaces['w']}}}id"] for reference in references] == ["1", "2"]
        texts = {note.attrib[f"{{{namespaces['w']}}}id"]: "".join(note.itertext()).strip() for note in notes}
        assert texts["1"] == "錄音.MP3 · 第 3 段 · 區塊 2–3；課堂.pdf · 第 4 頁 · 區塊 1"
        assert texts["2"] == "錄音.MP3 · 第 3 段 · 區塊 2"
        assert len(notes.findall(".//w:footnoteRef", namespaces)) == 2
        relationships = ET.fromstring(archive.read("word/_rels/document.xml.rels"))
        assert any(relation.attrib["Type"].endswith("/footnotes") and relation.attrib["Target"] == "footnotes.xml" for relation in relationships)
        content_types = ET.fromstring(archive.read("[Content_Types].xml"))
        assert any(part.attrib.get("PartName") == "/word/footnotes.xml" and part.attrib["ContentType"].endswith(".footnotes+xml") for part in content_types)
