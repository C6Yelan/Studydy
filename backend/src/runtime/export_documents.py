"""已發布資料的離線文件；不呼叫模型、不修改生成內容。"""
from functools import lru_cache
from datetime import datetime
from zoneinfo import ZoneInfo
from html import escape
from io import BytesIO
from pathlib import Path
import re
import zipfile

import pymupdf

from . import note_latex
from .note_content import CALLOUT_LABELS, body_html, numbered_heading, outline

FONT_DIR = Path('/usr/share/fonts/opentype/noto')
CSS = '''
@font-face { font-family: ExportCJK; src: url(NotoSansTC-Regular.otf); }
@font-face { font-family: ExportCJK; src: url(NotoSansTC-Bold.otf); font-weight: bold; }
body { font-family: ExportCJK; font-size: 12pt; line-height: 1.4; }
h1 { font-size: 18pt; } h2 { font-size: 14pt; margin: 8pt 0 4pt; }
p { margin: 0 0 6pt; white-space: pre-wrap; overflow-wrap: anywhere; }
table { border-collapse: collapse; width: 100%; margin: 0 0 6pt; }
td, th { border: .5pt solid #555; padding: 3pt; }
h3 { font-size: 13pt; margin: 6pt 0 3pt; }
ul { margin: 0 0 6pt; padding-left: 14pt; } li { margin: 0 0 2pt; }
.caption, .callout-title { font-weight: bold; margin: 6pt 0 3pt; }
.formula { text-align: center; margin: 6pt 0; }
.formula td { border: 0; padding: 0; }
.callout { background-color: #f3f7fc; border-left: 3pt solid #2f6bff; padding: 4pt 8pt; margin: 6pt 0; }
.outline-2 { margin-left: 16pt; }
'''
PAGE = pymupdf.paper_rect('a4')
CONTENT = PAGE + (51.024, 51.024, -51.024, -51.024)


@lru_cache(maxsize=2)
def font_bytes(weight):
    from fontTools.ttLib import TTCollection
    collection = TTCollection(FONT_DIR / f'NotoSansCJK-{weight}.ttc')
    font = next(font for font in collection.fonts if font['name'].getDebugName(6) == f'NotoSansCJKtc-{weight}')
    data = BytesIO()
    font.save(data)
    collection.close()
    return data.getvalue()


def font_archive():
    archive = pymupdf.Archive()
    for weight in ('Regular', 'Bold'):
        archive.add((font_bytes(weight), f'NotoSansTC-{weight}.otf'))
    return archive


def pdf(html, images=(), *, keep_headings=False, footnotes=None):
    archive = font_archive()
    for image in images:
        archive.add(image)
    forced_breaks, footer_height, footers = set(), 0, {}
    footer_css = CSS + 'body { font-size:10pt; } p { margin:0 0 3pt; }'
    while True:
        story = pymupdf.Story(html=html, user_css=CSS, archive=archive)
        if not keep_headings and not footnotes:
            break
        pairs = []
        node = story.body.first_child
        while node is not None:
            following = node.next
            if keep_headings and node.tagname in ('h1', 'h2', 'h3') and following is not None and following.tagname in ('p', 'ul', 'table', 'div'):
                index = len(pairs)
                heading_id, body_id = f'note-heading-{index}', f'note-body-{index}'
                node.set_id(heading_id)
                following.set_id(body_id)
                pairs.append((heading_id, body_id))
                if index in forced_breaks:
                    node.set_attribute('style', 'page-break-before:always;')
            node = following
        starts, markers = {}, {}
        def remember(position):
            if position.open_close & 1:
                starts[position.id] = position.page_num
                if footnotes and position.id in footnotes:
                    markers[position.id] = position.page_num
        content = CONTENT + (0, 0, 0, -footer_height)
        story.write(None, lambda _page, _filled: (PAGE, content, None), positionfn=remember)
        orphaned = {index for index, (heading, body) in enumerate(pairs) if starts[heading] != starts[body]}
        footers = {}
        for marker, page in markers.items():
            footers.setdefault(page, []).append(footnotes[marker])
        footers = {page: ''.join(f'<p>{escape(label)}</p>' for label in labels) for page, labels in footers.items()}
        needed = footer_height
        for footer in footers.values():
            measured = pymupdf.Story(html=footer, user_css=footer_css, archive=archive)
            _, filled = measured.place(pymupdf.Rect(0, 0, CONTENT.width, 10000))
            needed = max(needed, pymupdf.Rect(filled).height + 16)
        if not orphaned and needed == footer_height:
            story.reset()
            break
        forced_breaks.update(orphaned)
        footer_height = needed
    stream = BytesIO()
    writer = pymupdf.DocumentWriter(stream)
    more = True
    while more:
        device = writer.begin_page(PAGE)
        more, _ = story.place(CONTENT + (0, 0, 0, -footer_height))
        story.draw(device)
        writer.end_page()
    writer.close()
    with pymupdf.open(stream=stream.getvalue(), filetype='pdf') as document:
        for number, footer in footers.items():
            page = document[number - 1]
            top = CONTENT.y1 - footer_height
            page.draw_line((CONTENT.x0, top), (CONTENT.x1, top), color=(.6, .6, .6), width=.5)
            page.insert_htmlbox(pymupdf.Rect(CONTENT.x0, top + 8, CONTENT.x1, CONTENT.y1),
                                footer, css=footer_css, archive=archive, scale_low=1)
        stamp = datetime.now(ZoneInfo('Asia/Taipei')).strftime("D:%Y%m%d%H%M%S+08'00'")
        document.set_metadata({**document.metadata, 'creationDate':stamp, 'modDate':stamp})
        return document.tobytes(garbage=4, deflate=True)


def source_text(e):
    unit = '段' if re.search(r'\.(mp3|m4a|wav|mp4|youtube)$', e.get('source_name') or '', re.I) else '頁'
    position = f'第 {e.get("normalized_page") or e["page"]} {unit}'
    return f'{e.get("source_name") or (e.get("source_id") or "來源")[:8]} · {position}' + (
        f' · 區塊 {e["block_order"]}' if 'block_order' in e else '')


def md_text(value):
    return re.sub(r'([\\`*\[\]<>|])', r'\\\1', str(value))


def _note_title(view):
    return '重點整理' if view['kind'] == 'summary' else '完整筆記'


def document_metadata(view):
    # revision 帶有 knowledge-structure:sha256: 前綴；短碼取真正的 hash。
    return [f'教材：{view["material_name"]}（{view["material_id"][:8]}）',
            f'版本：{view["knowledge_structure_revision"].rsplit(":", 1)[-1][:8]}']


def _note_sources(evidence):
    """同檔同頁合併顯示，仍保留所有被引用的區塊，且不改保存的 Evidence。"""
    groups = {}
    for e in evidence:
        key = (e.get('source_id'), e.get('source_name'), e.get('normalized_page') or e['page'])
        sample, orders = groups.setdefault(key, (e, set()))
        if 'block_order' in e:
            orders.add(e['block_order'])
    for sample, orders in groups.values():
        label = source_text({k: v for k, v in sample.items() if k != 'block_order'})
        ranges = []
        for value in sorted(orders):
            if ranges and value == ranges[-1][1] + 1:
                ranges[-1][1] = value
            else:
                ranges.append([value, value])
        if ranges:
            label += ' · 區塊 ' + '、'.join(str(start) if start == end else f'{start}–{end}' for start, end in ranges)
        yield label


def _source_notice(view):
    status = view['source_status']
    return status['quality'] == 'needs_review' or status['processing'] != 'succeeded' or status['decision'] != 'retain'


def _md_items(items, indent=''):
    lines = []
    for item in items:
        lines.append(f'{indent}- {md_text(item["text"])}')
        lines.extend(_md_items(item.get('children') or [], indent + '  '))
    return lines


def _md_table(headers, rows):
    def cell(value):
        return md_text(value).replace('\n', '<br>')
    return ['| ' + ' | '.join(cell(c) for c in headers) + ' |', '| ' + ' | '.join('---' for _ in headers) + ' |',
            *('| ' + ' | '.join(cell(c) for c in row) + ' |' for row in rows), '']


def _md_blocks(blocks):
    lines = []
    for block in blocks:
        kind = block['type']
        if kind == 'text':
            lines.extend([md_text(block['text']), ''])
        elif kind == 'bullets':
            lines.extend([*_md_items(block['items']), ''])
        elif kind == 'table':
            if block['caption']:
                lines.extend([f'**{md_text(block["caption"])}**', ''])
            lines.extend(_md_table(block['headers'], block['rows']))
        elif kind == 'formula':
            # 公式以 LaTeX 原樣輸出，不做 Markdown 跳脫。
            lines.extend(['$$', block['latex'], '$$', ''])
        else:
            title = CALLOUT_LABELS[block['style']] + (f'：{block["title"]}' if block['title'] else '')
            lines.extend([f'> **{md_text(title)}**', '>', *(f'> {line}' for line in _md_items(block['items'])), ''])
    return lines


def note_markdown(view):
    lines = [f'# {_note_title(view)}', '', *(md_text(line) for line in document_metadata(view)), '']
    if _source_notice(view):
        lines += ['> 教材仍有待核對或尚未完整處理的內容，請配合段落來源閱讀。', '']
    numbers = {o['paragraph_id']: o for o in outline(view['paragraphs'])}
    anchors, used = {}, set()
    if view['paragraphs']:
        lines.extend(['## 章節大綱', ''])
        for p in view['paragraphs']:
            number = numbers.get(p['id'])
            heading = numbered_heading(number) if number else p['heading']
            base = re.sub(r'[^\w-]+', '-', heading.casefold()).strip('-') or 'section'
            anchor, suffix = base, 2
            while anchor in used:
                anchor = f'{base}-{suffix}'
                suffix += 1
            used.add(anchor)
            anchors[p['id']] = anchor
            indent = '  ' if number and number['level'] == 2 else ''
            lines.append(f'{indent}- [{md_text(heading)}](#{anchor})')
        lines.append('')

    def paragraph(p, number=None):
        if 'blocks' not in p:
            # 舊版段落格式：保留可讀，直到重新生成。
            lines.extend([f'## {md_text(p["heading"])}', '', md_text(p['text']), ''])
            if p.get('formula'):
                lines.extend([md_text(p['formula']), ''])
            if p.get('table'):
                lines.extend(_md_table(p['table']['headers'], p['table']['rows']))
            depth = 2
        else:
            depth = 1 + (number['level'] if number else p['level'])
            heading = numbered_heading(number) if number else p['heading']
            lines.extend([f'{"#" * depth} {md_text(heading)}', '', *_md_blocks(p['blocks'])])
        lines.extend(f'來源：{md_text(label)}  ' for label in _note_sources(p['evidence']))
        lines.append('')
        return depth

    for p in view['paragraphs']:
        lines.extend([f'<a id="{anchors[p["id"]]}"></a>', ''])
        depth = paragraph(p, numbers.get(p['id']))
        for a in view['annotations']:
            if a['paragraph_id'] == p['id']:
                lines.extend([f'{"#" * (depth + 1)} 個人註記', '', md_text(a['body']), ''])
    orphaned = [a for a in view['annotations'] if a['paragraph_id'] is None]
    if orphaned:
        lines.extend(['# 先前段落的註記', '', '以下註記未能精確對應目前段落，保留當時內容。', ''])
        for a in orphaned:
            lines.extend([f'原版本：{a["original_revision"].rsplit(":", 1)[-1][:8]}', ''])
            depth = paragraph(a['original_paragraph'])
            lines.extend([f'{"#" * (depth + 1)} 個人註記', '', md_text(a['body']), ''])
    return '\n'.join(lines)


def note_pdf(view):
    images, footnotes = [], {}
    parts = [f'<h1>{_note_title(view)}</h1>', '<p>' + escape('\n'.join(document_metadata(view))) + '</p>']
    numbers = {o['paragraph_id']: o for o in outline(view['paragraphs'])}
    if view['paragraphs']:
        parts.append('<h2>章節大綱</h2><div>' + ''.join(
            f'<p>{escape(numbered_heading(numbers[p["id"]]) if p["id"] in numbers else p["heading"])}</p>'
            for p in view['paragraphs']) + '</div>')
    for p in view['paragraphs']:
        heading = numbered_heading(numbers[p['id']]) if p['id'] in numbers else p['heading']
        parts.append(body_html([{**p, 'heading':heading}], images, numbered=False))
        if p['evidence']:
            number = len(footnotes) + 1
            marker = f'source-{number}'
            footnotes[marker] = f'[{number}] ' + '；'.join(_note_sources(p['evidence']))
            parts.append(f'<p id="{marker}">[{number}]</p>')
        parts.extend(f'<p>個人註記：{escape(a["body"])}</p>' for a in view['annotations'] if a['paragraph_id'] == p['id'])
    orphaned = [a for a in view['annotations'] if a['paragraph_id'] is None]
    if orphaned:
        parts.append('<h1>先前段落的註記</h1><p>以下註記未能精確對應目前段落，保留當時內容。</p>')
        for a in orphaned:
            p = a['original_paragraph']
            parts.append(f'<p>原版本：{escape(a["original_revision"].rsplit(":", 1)[-1][:8])}</p>' + body_html([p], images, numbered=False))
            if p['evidence']:
                number = len(footnotes) + 1
                marker = f'source-{number}'
                footnotes[marker] = f'[{number}] ' + '；'.join(_note_sources(p['evidence']))
                parts.append(f'<p id="{marker}">[{number}]</p>')
            parts.append(f'<p>個人註記：{escape(a["body"])}</p>')
    return pdf(''.join(parts), images, keep_headings=True, footnotes=footnotes)


def xml_text(value):
    # XML 1.0 允許 tab／換行；移除 Office 無法保存的控制字元與孤立 surrogate。
    return re.sub('[\x00-\x08\x0b\x0c\x0e-\x1f\ud800-\udfff\ufffe\uffff]', '', value)


def _docx_items(container, items, depth=1):
    for item in items:
        container.add_paragraph(xml_text(item['text']), style='List Bullet' if depth == 1 else f'List Bullet {depth}')
        _docx_items(container, item.get('children') or [], depth + 1)


def _docx_table(doc, headers, rows):
    table = doc.add_table(rows=1, cols=len(headers))
    table.style = 'Table Grid'
    for cell, value in zip(table.rows[0].cells, headers):
        cell.text = xml_text(value)
    for row in rows:
        for cell, value in zip(table.add_row().cells, row):
            cell.text = xml_text(value)
    # 相鄰表格之間須隔一段，否則 Word 會把它們合成同一個表格。
    doc.add_paragraph()


def _docx_blocks(doc, blocks):
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    from docx.shared import Pt
    for block in blocks:
        kind = block['type']
        if kind == 'text':
            doc.add_paragraph(xml_text(block['text']))
        elif kind == 'bullets':
            _docx_items(doc, block['items'])
        elif kind == 'table':
            if block['caption']:
                doc.add_paragraph().add_run(xml_text(block['caption'])).bold = True
            _docx_table(doc, block['headers'], block['rows'])
        elif kind == 'formula':
            # Word 不能直接顯示 LaTeX；嵌入與 PDF 相同的排版圖，替代文字保留 LaTeX 原式。
            data, width, _ = note_latex.render_png(block['latex'])
            paragraph = doc.add_paragraph()
            paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
            section = doc.sections[-1]
            page_width = section.page_width - section.left_margin - section.right_margin
            paragraph.add_run().add_picture(BytesIO(data), width=min(Pt(width), page_width))
            paragraph._p.xpath('.//wp:docPr')[0].set('descr', xml_text(block['latex']))
        else:
            box = doc.add_table(rows=1, cols=1)
            box.style = 'Table Grid'
            cell = box.rows[0].cells[0]
            shading = OxmlElement('w:shd')
            shading.set(qn('w:val'), 'clear')
            shading.set(qn('w:fill'), 'F3F7FC')
            cell._tc.get_or_add_tcPr().append(shading)
            title = CALLOUT_LABELS[block['style']] + (f'：{block["title"]}' if block['title'] else '')
            cell.paragraphs[0].add_run(xml_text(title)).bold = True
            _docx_items(cell, block['items'])
            doc.add_paragraph()


def note_docx(view):
    from docx import Document
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    from docx.shared import Pt
    doc = Document()
    doc.core_properties.author = 'Studydy'
    doc.core_properties.last_modified_by = 'Studydy'
    doc.core_properties.comments = ''
    for name in ('Normal', 'Title', 'Heading 1', 'Heading 2', 'Heading 3'):
        style = doc.styles[name]
        style.font.name = 'Noto Sans CJK TC'
        fonts = style.element.get_or_add_rPr().find(qn('w:rFonts'))
        if fonts is None:
            fonts = OxmlElement('w:rFonts')
            style.element.get_or_add_rPr().append(fonts)
        fonts.set(qn('w:eastAsia'), 'Noto Sans CJK TC')
    doc.add_heading(_note_title(view), 0)
    doc.add_paragraph(xml_text('\n'.join(document_metadata(view))))
    if _source_notice(view):
        doc.add_paragraph('教材仍有待核對或尚未完整處理的內容，請配合段落來源閱讀。')
    numbers = {o['paragraph_id']: o for o in outline(view['paragraphs'])}
    if view['paragraphs']:
        doc.add_heading('章節大綱', 1)
        for p in view['paragraphs']:
            number = numbers.get(p['id'])
            item = doc.add_paragraph(xml_text(numbered_heading(number) if number else p['heading']))
            if number and number['level'] == 2:
                item.paragraph_format.left_indent = Pt(18)


    def paragraph(p, number=None):
        if 'blocks' not in p:
            # 舊版段落格式：保留可讀，直到重新生成。
            doc.add_heading(xml_text(p['heading']), 1)
            doc.add_paragraph(xml_text(p['text']))
            if p.get('formula'):
                doc.add_paragraph(xml_text(p['formula']))
            if p.get('table'):
                _docx_table(doc, p['table']['headers'], p['table']['rows'])
            level = 1
        else:
            level = number['level'] if number else p['level']
            doc.add_heading(xml_text(numbered_heading(number) if number else p['heading']), level)
            _docx_blocks(doc, p['blocks'])
        for e in p['evidence']:
            doc.add_paragraph(xml_text('來源：' + source_text(e)))
        return level

    for p in view['paragraphs']:
        level = paragraph(p, numbers.get(p['id']))
        for a in view['annotations']:
            if a['paragraph_id'] == p['id']:
                doc.add_heading('個人註記', level + 1)
                doc.add_paragraph(xml_text(a['body']))
    orphaned = [a for a in view['annotations'] if a['paragraph_id'] is None]
    if orphaned:
        doc.add_heading('先前段落的註記', 1)
        for a in orphaned:
            doc.add_paragraph(xml_text('原版本：' + a['original_revision'].rsplit(':', 1)[-1][:8]))
            level = paragraph(a['original_paragraph'])
            doc.add_heading('個人註記', level + 1)
            doc.add_paragraph(xml_text(a['body']))
    stream = BytesIO()
    doc.save(stream)
    from xml.etree import ElementTree as ET
    with zipfile.ZipFile(BytesIO(stream.getvalue())) as archive:
        files = []
        for item in archive.infolist():
            content = archive.read(item.filename)
            if item.filename == 'docProps/core.xml':
                root = ET.fromstring(content)
                for tag in ('created', 'modified'):
                    root.find('{http://purl.org/dc/terms/}' + tag).text = datetime.now(ZoneInfo('Asia/Taipei')).isoformat(timespec='seconds')
                content = ET.tostring(root, encoding='utf-8', xml_declaration=True)
            files.append((item.filename, content))
    return zip_files(files)


def zip_files(files):
    stream = BytesIO()
    with zipfile.ZipFile(stream, 'w', zipfile.ZIP_DEFLATED) as archive:
        for name, content in files:
            entry = zipfile.ZipInfo(name, datetime.now(ZoneInfo('Asia/Taipei')).timetuple()[:6])
            archive.writestr(entry, content, compress_type=zipfile.ZIP_DEFLATED)
    return stream.getvalue()


def diagnosis_pdf(view):
    counts = {key: sum(len(item['points']) for item in view[key]) for key in ('strengths', 'weaknesses')}
    head = f'<h1>教材學習診斷</h1><p>優勢重點：{counts["strengths"]}　待加強重點：{counts["weaknesses"]}</p>'
    recommendations = view['recommendations']
    points = [(label, finding, point) for key, label in [('weaknesses', '弱點'), ('strengths', '優勢')]
              for finding in view[key] for point in finding['points']]
    selected_recommendations, selected = [], []

    def render():
        omitted = []
        if len(points) > len(selected):
            omitted.append(f'{len(points) - len(selected)} 個診斷重點')
        if len(recommendations) > len(selected_recommendations):
            omitted.append(f'{len(recommendations) - len(selected_recommendations)} 項複習建議')
        footer = f'<p>另有 {"、".join(omitted)}未列於摘要；完整內容請回到 Studydy 查看。</p>' if omitted else ''
        title = '<h2>建議先複習</h2>' if selected_recommendations else ''
        return pdf(head + title + ''.join(selected_recommendations) + ''.join(selected) + footer)

    def fits():
        with pymupdf.open(stream=render(), filetype='pdf') as doc:
            return len(doc) == 1

    for r in recommendations:
        selected_recommendations.append(f'<p>{escape(r["label"])}：{escape(r["reason"])}</p>')
        if not fits():
            selected_recommendations.pop()
            break
    for label, finding, point in points:
        selected.append(f'<h2>{label}：{escape(finding["label"])}</h2><p>{escape(point["title"])}\n答對 {point["correct_count"]} / {point["answer_count"]} 次</p>' +
            ''.join(f'<p>{escape(source_text(e))}</p>' for e in point['evidence']))
        if not fits():
            selected.pop()
            break
    return render()


def decode_png(value):
    from base64 import b64decode
    from struct import unpack
    from .errors import SourceError
    try:
        data = b64decode(value, validate=True)
        if not data.startswith(b'\x89PNG\r\n\x1a\n'):
            raise ValueError()
        width, height = unpack('>II', data[16:24])
        if not 0 < width * height <= 16000000:
            raise ValueError()
        with pymupdf.open(stream=data, filetype='png') as document:
            document[0].get_pixmap()
        return data
    except Exception as error:
        raise SourceError('REQUEST_INVALID') from error


def compact_pdf(data):
    with pymupdf.open(stream=data, filetype='pdf') as doc:
        # MuPDF 的原生 CFF 子集化對 Noto TC 會出現 Index bounds；採已安裝的 fontTools 實作。
        doc.subset_fonts(fallback=True)
        return doc.tobytes(garbage=4, deflate=True)
